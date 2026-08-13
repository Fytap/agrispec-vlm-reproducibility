#!/usr/bin/env python3
"""Create the editable SAT Declaration of Interest document."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "Bounded Instance-Support Evaluation for Finite Crop--Weed Review Queues"
AUTHORS = "Shuhao Liu; Zhuo Chen; Zhi Ling; Yu Yan; Jie Liu; Qiuxue Wu; Ziyi Kuang"
DECLARATION = (
    "The authors declare that they have no known competing financial interests "
    "or personal relationships that could have appeared to influence the work "
    "reported in this paper."
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def build(output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.12

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(16)
    run = heading.add_run("Declaration of Competing Interests")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(36, 49, 58)

    table = doc.add_table(rows=2, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "E7EFF6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Manuscript")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(63, 111, 168)
    p = cell.add_paragraph(TITLE)
    p.paragraph_format.space_after = Pt(5)
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(11)

    cell = table.cell(1, 0)
    p = cell.paragraphs[0]
    r = p.add_run("Authors")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(63, 111, 168)
    p = cell.add_paragraph(AUTHORS)
    p.paragraph_format.space_after = Pt(5)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    statement = doc.add_paragraph(DECLARATION)
    statement.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    statement.paragraph_format.space_before = Pt(8)
    statement.paragraph_format.space_after = Pt(18)
    statement.runs[0].font.size = Pt(11)

    corr = doc.add_paragraph()
    corr.add_run("Corresponding author: ").bold = True
    corr.add_run("Zhuo Chen (zhuoc@chalmers.se)")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
