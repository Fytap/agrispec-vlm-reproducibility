#!/usr/bin/env python3
"""Create the standalone informed-consent statement for SAT submission."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


MANUSCRIPT_TITLE = (
    "Bounded Instance-Support Evaluation for Finite Crop--Weed Review Queues"
)
STATEMENT = (
    "All annotators provided informed consent prior to participation in the study."
)


def set_run_font(
    run,
    *,
    size: float,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor = RGBColor(0, 0, 0),
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def build(output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    doc.core_properties.title = "Ethics and Informed Consent Statement"
    doc.core_properties.subject = MANUSCRIPT_TITLE

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(12)
    set_run_font(title.add_run("Ethics and Informed Consent Statement"), size=20, bold=True)

    manuscript_label = doc.add_paragraph()
    manuscript_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    manuscript_label.paragraph_format.space_after = Pt(4)
    set_run_font(
        manuscript_label.add_run("Manuscript"),
        size=10,
        bold=True,
        color=RGBColor(46, 116, 181),
    )

    manuscript = doc.add_paragraph()
    manuscript.alignment = WD_ALIGN_PARAGRAPH.CENTER
    manuscript.paragraph_format.space_after = Pt(28)
    set_run_font(manuscript.add_run(MANUSCRIPT_TITLE), size=11, italic=True)

    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    statement.paragraph_format.left_indent = Inches(0.35)
    statement.paragraph_format.right_indent = Inches(0.35)
    statement.paragraph_format.space_before = Pt(8)
    statement.paragraph_format.space_after = Pt(8)
    statement.paragraph_format.line_spacing = 1.20
    set_run_font(statement.add_run(STATEMENT), size=12)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
